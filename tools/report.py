"""PDF/Excel snapshot reports and a safe consolidated export ZIP.

PDF labels currently require Latin/English text. Unsupported scripts fail
explicitly rather than rendering broken glyphs. Excel retains Unicode text.
BOQ formulas carry cached snapshot results and request recalculation on open;
editing the workbook does not revise the app's approved project baseline.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path, PurePosixPath
import reportlab
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from config import DISCLAIMER
from schemas import Contract, LayoutAlternative, Requirements, content_hash
from tools.cost import CostReport, dec, money
from tools.labour import LabourReport
from tools.schedule_4d import ScheduleResult
from tools.twin import TwinSnapshot
from agents.structural_agent import StructuralReport
from agents.mep_agent import MEPReport


class ReportInput(Contract):
    project_name: str
    report_date: date
    requirements: Requirements
    layout: LayoutAlternative
    reviewed_layout_hash: str
    cost: CostReport | None = None
    cost_hash: str | None = None
    schedule: ScheduleResult | None = None
    schedule_hash: str | None = None
    labour: LabourReport | None = None
    labour_hash: str | None = None
    monitoring: TwinSnapshot | None = None
    monitoring_hash: str | None = None
    structural: StructuralReport | None = None
    structural_hash: str | None = None
    mep: MEPReport | None = None
    mep_hash: str | None = None


@dataclass(frozen=True)
class ExportAttachment:
    name: str
    content: bytes


@dataclass(frozen=True)
class ReportFiles:
    pdf_bytes: bytes
    xlsx_bytes: bytes
    zip_bytes: bytes
    input_hash: str


class ReportExportError(ValueError):
    """Export limitation or validation failure safe to show in the UI."""


def validate_report(request: ReportInput) -> ReportInput:
    request=ReportInput.model_validate(request.model_dump(mode='json'))
    if request.reviewed_layout_hash!=content_hash(request.layout):
        raise ReportExportError('Layout snapshot changed. Review the current layout before reporting.')
    if {f.floor_id for f in request.layout.floors}!={f.id for f in request.requirements.floors}:
        raise ReportExportError('Report layout does not match the requested floor IDs.')
    for name in ('cost','schedule','labour','monitoring','structural','mep'):
        value=getattr(request,name);fingerprint=getattr(request,name+'_hash')
        if (value is None)!=(fingerprint is None) or (value is not None and fingerprint!=content_hash(value)):
            raise ReportExportError('A report source snapshot is missing or stale.')
    if request.labour and request.schedule and request.labour.schedule_hash!=content_hash(request.schedule):
        raise ReportExportError('Labour source does not match the supplied schedule.')
    return request


def _display(value) -> str:
    if value is None: return 'Unresolved'
    if isinstance(value,float): return f'{value:,.2f}'
    return str(value)


def _latin(text: str) -> str:
    for a,b in [('–','-'),('—','-'),('‑','-'),('−','-'),('×','x'),('²','2'),('³','3'),('’',"'"),('‘',"'"),('“','"'),('”','"')]:
        text=text.replace(a,b)
    try: text.encode('cp1252')
    except UnicodeEncodeError:
        raise ReportExportError('PDF export currently requires English/Latin labels. Translate unsupported script before export; Unicode remains supported in Excel.') from None
    return text


def _metrics(request: ReportInput) -> list[tuple[str,object]]:
    covered=sum(f.covered_area_m2 for f in request.layout.floors)
    result=[('Project',request.project_name),('Report date',request.report_date.isoformat()),
        ('City / authority',request.requirements.city+' / '+request.requirements.authority),
        ('Layout',request.layout.id),('Requested quality',request.requirements.quality.value),
        ('Covered envelope area (m2)',covered),('Covered envelope area (sq ft)',covered/(.3048**2))]
    if request.cost:
        c=request.cost
        result.extend([('Estimate basis',c.selected_basis),('Construction (PKR)',c.construction_pkr),
            ('Extras (PKR)',c.extras_pkr),('Contingency (PKR)',c.contingency_pkr),('Tax scenario (PKR)',c.tax_pkr),
            ('Grand total (PKR)',c.grand_total_pkr),('Known BOQ subtotal (PKR)',c.boq_known_subtotal_pkr)])
    if request.schedule: result.append(('Schedule duration (calendar days)',request.schedule.duration_calendar_days))
    if request.labour: result.extend([('Peak planned workers',request.labour.peak_workers),('Paid person-days',request.labour.paid_person_days),('Wages (PKR)',request.labour.total_wages_pkr)])
    return result


def render_pdf(request: ReportInput) -> bytes:
    request=validate_report(request)
    font_root=Path(reportlab.__file__).parent/'fonts'
    for name,filename in [('GharBuilderSans','Vera.ttf'),('GharBuilderBold','VeraBd.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():
            font_path=font_root/filename
            if not font_path.exists():
                raise ReportExportError('The ReportLab bundled Vera fonts are missing; reinstall ReportLab before PDF export.')
            pdfmetrics.registerFont(TTFont(name,str(font_path)))
    styles=getSampleStyleSheet()
    styles['Title'].fontName='GharBuilderBold'
    body=ParagraphStyle('ReportBody',parent=styles['BodyText'],fontName='GharBuilderSans',fontSize=8.5,
                        leading=11,spaceAfter=4,splitLongWords=True)
    heading=ParagraphStyle('ReportHeading',parent=styles['Heading2'],fontName='GharBuilderBold',fontSize=12,textColor=colors.HexColor('#3949AB'),spaceBefore=9,spaceAfter=6)
    def p(value): return Paragraph(escape(_latin(_display(value))).replace('\n','<br/>'),body)
    story=[]
    story.append(Paragraph('GHARBUILDER AI',styles['Title']))
    story.append(p('House concept and estimation report'))
    story.append(p(DISCLAIMER))
    def section(title,headers,rows,widths):
        story.append(Paragraph(_latin(title),heading))
        if not rows: story.append(p('Not supplied.'));return
        cells=[[p(v) for v in headers]]+[[p(v) for v in row] for row in rows]
        table=LongTable(cells,colWidths=widths,repeatRows=1,hAlign='LEFT',splitInRow=1)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E9EBFA')),
            ('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F7FB')]),
            ('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),
            ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),
            ('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#BFC7DB'))]))
        story.append(table)
    section('Project overview',['Metric','Value'],_metrics(request),[210,313])
    section('Floor program',['Floor','Elevation m','Zone rooms','Covered m2'],
        [[f.floor_id,f.elevation_m,len(f.rooms),f.covered_area_m2] for f in request.layout.floors],[170,100,110,143])
    brief=request.requirements
    section('Confirmed brief',['Requirement','Value'],[
        ['Plot width / depth m',f'{brief.plot.width_m:g} / {brief.plot.depth_m:g}'],
        ['Facing / corner',brief.plot.facing+' / '+str(brief.plot.corner)],
        ['Marla conversion sq ft',brief.plot.marla_sqft],['Family / cars',f'{brief.family_size} / {brief.car_spaces}'],
        ['Stairs / lift',brief.stairs_position+' / '+str(brief.lift)],['Style',brief.style],
        ['Accessibility',', '.join(brief.accessibility_needs) or 'None recorded'],
        ['Green features',', '.join(brief.green_features) or 'None recorded'],
        ['Requested start',brief.start_date.isoformat()],['Budget PKR',brief.budget_pkr]], [210,313])
    if request.cost:
        c=request.cost
        source_rows=[]
        for label,e in [('Area benchmark',c.area_benchmark_evidence),('City multiplier',c.city_factor_evidence)]:
            source_rows.append([label,(e.source+' / '+(e.effective_date.isoformat() if e.effective_date else 'Date missing')+
                (' / verified record' if e.verified else ' / unverified')) if e else 'Not supplied'])
        section('Benchmark evidence',['Input','Source / date'],source_rows,[150,373])
        section('BOQ and explicit extras',['Description / floor','Quantity / unit','Rate PKR','Amount PKR'],
            [[r.description+(' / '+r.floor_id if r.floor_id else ''),f'{r.quantity:,.3f} {r.unit}',r.effective_rate_pkr,r.amount_pkr] for r in c.rows],
            [220,110,90,103])
        section('Rate and quantity evidence',['Item','Rate date / verification','Source / formula'],
            [[r.id,(r.rate_evidence.effective_date.isoformat() if r.rate_evidence and r.rate_evidence.effective_date else 'Date missing')+
               (' / verified record' if r.rate_evidence and r.rate_evidence.verified else ' / unverified'),
               (r.rate_evidence.source if r.rate_evidence else 'Missing quotation')+'\n'+(r.rate_evidence.source_url or '' if r.rate_evidence else '')+'\n'+r.formula+'\nQuantity: '+r.source] for r in c.rows],
            [105,115,303])
        section('Milestone scenario',['Milestone','Percent','Amount PKR'],
            [[x['label'],x['percent'],x['amount_pkr']] for x in c.milestone_amounts_pkr],[290,100,133])
    if request.schedule:
        section('Construction schedule',['Task','Start','Finish exclusive','Release exclusive'],
            [[t.activity.name,t.start.isoformat(),t.finish_exclusive.isoformat(),t.release_exclusive.isoformat()] for t in request.schedule.tasks],
            [238,95,95,95])
    if request.labour:
        section('Labour allowances',['Trade','Paid person-days','Peak workers','Wages PKR'],
            [[t.trade,t.paid_person_days,t.peak_workers,t.wages_pkr] for t in request.labour.totals],[205,110,100,108])
    if request.structural:
        s=request.structural
        section('Preliminary structural coordination',['Metric','Value'],[
            ['Status',s.status],['Candidate grid points',len(s.grid)],['Foundation',s.foundation_status],
            ['Seismic record',s.seismic_status]],[210,313])
    if request.mep:
        m=request.mep
        section('Indicative MEP allowances',['Metric','Value'],[
            ['Connected electrical kW',m.total_connected_kw],['Indicative peak demand kW',m.indicative_peak_demand_kw],
            ['Daily energy allowance kWh',m.estimated_daily_kwh],['Water storage litres',m.total_storage_litres],
            ['PV energy target kWp',m.solar_energy_target_kwp],['Area-limited PV kWp',m.solar_area_limited_kwp]],[290,233])
    if request.monitoring:
        m=request.monitoring
        section('User-reported construction monitoring',['Metric','Value'],[
            ['As of',m.as_of.isoformat()],['Planned productive effort %',m.planned_effort_percent],
            ['Reported actual effort-weighted %',m.actual_reported_effort_weighted_percent],
            ['Reporting coverage %',m.progress_reporting_coverage_percent],
            ['Known actual spend subtotal PKR',m.known_actual_cost_subtotal_pkr],
            ['Actual spend total PKR',m.actual_cost_total_pkr]],[290,233])
    story.append(Paragraph('Assumptions and review items',heading))
    notes=[*request.layout.unresolved_issues,
        'Room dimensions and areas are centre-line concept zones, not certified clear internal dimensions.',
        'Optional source stages omitted from this report are not represented as completed work.',
        'Source hashes identify the supplied snapshots. Upstream layout/source consistency also requires the app revision controls.']
    for value in (request.cost,request.schedule,request.labour):
        if value: notes+=value.warnings
    if request.monitoring: notes+=request.monitoring.alerts+request.monitoring.limitations
    if request.structural: notes+=request.structural.warnings
    if request.mep: notes+=request.mep.unresolved
    for note in dict.fromkeys(notes): story.append(p(note))
    def page(canvas,doc):
        canvas.saveState();canvas.setFont('GharBuilderSans',7);canvas.setFillColor(colors.HexColor('#536176'))
        canvas.drawString(36,22,'GharBuilder AI | preliminary concept | '+request.report_date.isoformat())
        canvas.drawRightString(A4[0]-36,22,'Page '+str(doc.page));canvas.restoreState()
    output=io.BytesIO()
    doc=SimpleDocTemplate(output,pagesize=A4,rightMargin=36,leftMargin=36,topMargin=32,bottomMargin=38,
                          title='GharBuilder concept and estimation report',author='GharBuilder AI')
    doc.build(story,onFirstPage=page,onLaterPages=page)
    return output.getvalue()


def _text(value):
    """Prevent spreadsheet formula injection in user/source text."""
    if isinstance(value,str) and value.startswith(('=','+','-','@')): return "'"+value
    return value


def render_excel(request: ReportInput) -> bytes:
    request=validate_report(request)
    wb=Workbook();summary=wb.active;summary.title='Overview'
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    summary.append(['GHARBUILDER AI','Report snapshot'])
    summary.append(['Metric','Value'])
    for key,value in _metrics(request): summary.append([key,_text(value)])
    summary.append(['Editing','Overview is an approved snapshot. Re-run the app to revise it; BOQ formulas update detail only.'])
    boq=wb.create_sheet('BOQ')
    boq.append(['Item ID','Description','Floor','Unit','Quantity','Effective rate PKR','Amount PKR','Phase','Category','Quantity formula','Quantity source'])
    cache={}
    for index,row in enumerate(request.cost.rows if request.cost else [],2):
        formula=f'=IF(E{index}=0,0,IF(ISNUMBER(F{index}),ROUND(E{index}*F{index},2),"Unpriced"))'
        boq.append([_text(row.id),_text(row.description),_text(row.floor_id),_text(row.unit),row.quantity,
            row.effective_rate_pkr,formula,row.phase,row.category,_text(row.formula),_text(row.source)])
        cache[f'G{index}']=row.amount_pkr if row.amount_pkr is not None else 'Unpriced'
    last=boq.max_row
    if last>=2:
        boq.append(['Known subtotal (all rows, including extras)','','','','','',f'=SUM(G2:G{last})'])
        cache[f'G{last+1}']=money(sum((dec(r.amount_pkr) for r in request.cost.rows if r.amount_pkr is not None),dec(0)))
    sources=wb.create_sheet('Sources')
    sources.append(['Item ID','Rate ID','Source','Effective date','Source URL','Verified record','Base rate PKR'])
    for row in request.cost.rows if request.cost else []:
        e=row.rate_evidence
        sources.append([_text(row.id),row.rate_id,_text(e.source) if e else 'Missing quotation',e.effective_date if e else None,
                        _text(e.source_url) if e else None,e.verified if e else False,row.base_rate_pkr])
    schedule=wb.create_sheet('Schedule')
    schedule.append(['Activity ID','Task','Floor','Start','Finish exclusive','Release exclusive','Productive days','Crew'])
    for task in request.schedule.tasks if request.schedule else []:
        schedule.append([task.activity.id,_text(task.activity.name),task.activity.floor_id,task.start,task.finish_exclusive,
                         task.release_exclusive,task.activity.duration_days,task.activity.crew_size])
    labour=wb.create_sheet('Labour')
    labour.append(['Trade','Paid person-days','Productive person-days','Peak workers','Daily wage PKR','Wages PKR'])
    for row in request.labour.totals if request.labour else []:
        labour.append([row.trade,row.paid_person_days,row.productive_person_days,row.peak_workers,row.daily_wage_pkr,row.wages_pkr])
    notes=wb.create_sheet('Notes');notes.append(['Assumptions and limitations'])
    note_values=[DISCLAIMER,'Blank prices are unresolved; a known subtotal is not a complete project estimate.',
                 'The BOQ contains supplied quantity allowances; formula editing does not change the app baseline.',
                 'Recalculation is requested on open. Cached BOQ values are calculated from the exported snapshot in Python.',
                 'Floor/room costs, if present in the PDF, are area allocations rather than measured room BOQs.',
                 *request.layout.unresolved_issues]
    for value in (request.cost,request.schedule,request.labour):
        if value: note_values+=value.warnings
    for note in dict.fromkeys(note_values): notes.append([_text(note)])
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='3949AB');cell.font=Font(name='Calibri',bold=True,color='FFFFFF',size=11)
        ws.row_dimensions[1].height=28
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                cell.font=Font(name='Calibri',size=10,color='1F3249')
                if isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool): cell.number_format='#,##0.00'
                if isinstance(cell.value,date): cell.number_format='yyyy-mm-dd'
        for column in ws.columns:
            letter=column[0].column_letter
            ws.column_dimensions[letter].width=24 if ws.title!='Notes' else 110
        if ws.title in {'BOQ','Sources','Schedule'}: ws.column_dimensions['B'].width=42
        if ws.title=='Overview': ws.column_dimensions['A'].width=38;ws.column_dimensions['B'].width=72
        if ws.title=='Sources': ws.column_dimensions['C'].width=55;ws.column_dimensions['E'].width=55
        for row in ws.iter_rows(min_row=2):
            # Estimate wrapped height for export usability; Excel can still autofit.
            lines=max(max((len(str(c.value or ''))//max(8,int(ws.column_dimensions[c.column_letter].width))+1),str(c.value or '').count('\n')+1) for c in row)
            ws.row_dimensions[row[0].row].height=max(26,min(180,15*lines+6))
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:1'
    output=io.BytesIO();wb.save(output)
    # Add accurate cached snapshot values to formula cells without removing formulas.
    source=zipfile.ZipFile(io.BytesIO(output.getvalue()));result=io.BytesIO()
    namespace='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('',namespace)
    with zipfile.ZipFile(result,'w',zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data=source.read(item.filename)
            if item.filename=='xl/worksheets/sheet2.xml':
                tree=ET.fromstring(data)
                for cell in tree.findall('.//{'+namespace+'}c'):
                    if cell.get('r') not in cache: continue
                    value=cache[cell.get('r')]
                    v=cell.find('{'+namespace+'}v')
                    if v is None: v=ET.SubElement(cell,'{'+namespace+'}v')
                    if isinstance(value,str): cell.set('t','str');v.text=value
                    else: cell.attrib.pop('t',None);v.text=str(value)
                data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
            target.writestr(item,data)
    source.close()
    return result.getvalue()


def report_fingerprint(request: ReportInput, attachments: list[ExportAttachment]) -> str:
    return content_hash({'report':request.model_dump(mode='json'),'attachments':[
        {'name':a.name,'sha256':hashlib.sha256(a.content).hexdigest()} for a in attachments]})


def build_reports(request: ReportInput, attachments: list[ExportAttachment] | None=None) -> ReportFiles:
    request=validate_report(request);attachments=attachments or []
    pdf=render_pdf(request);excel=render_excel(request)
    allowed={'.png','.dxf','.ifc','.glb','.obj','.json','.txt'}
    entries={};total=0
    def add(name,content):
        nonlocal total
        path=PurePosixPath(name)
        if str(path)!=name or path.is_absolute() or '..' in path.parts or '\\' in name or not re.fullmatch(r'[A-Za-z0-9_./-]+',name):
            raise ReportExportError('Unsafe export attachment path.')
        if path.suffix.lower() not in allowed or name in entries:
            raise ReportExportError('Unsupported or duplicate geometry attachment.')
        total+=len(content)
        if total>200*1024*1024: raise ReportExportError('Geometry attachments exceed the 200 MiB uncompressed limit.')
        entries[name]=content
    for attachment in attachments:
        if len(attachment.content)>100*1024*1024: raise ReportExportError('An attachment exceeds 100 MiB.')
        if PurePosixPath(attachment.name).suffix.lower()=='.zip':
            with zipfile.ZipFile(io.BytesIO(attachment.content)) as archive:
                if sum(i.file_size for i in archive.infolist())>200*1024*1024-total:
                    raise ReportExportError('Expanded geometry archive exceeds the size limit.')
                for item in archive.infolist():
                    if not item.is_dir(): add('geometry/'+item.filename,archive.read(item))
        else: add('geometry/'+attachment.name,attachment.content)
    output=io.BytesIO();fingerprint=report_fingerprint(request,attachments)
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('GharBuilder_Report.pdf',pdf);archive.writestr('GharBuilder_BOQ.xlsx',excel)
        archive.writestr('report_manifest.json',json.dumps({'input_hash':fingerprint,'report_date':request.report_date.isoformat(),
            'geometry_files':list(entries),'geometry_status':'Included supplied exports' if entries else 'No geometry exports supplied',
            'disclaimer':DISCLAIMER},indent=2))
        for name,data in entries.items(): archive.writestr(name,data)
    return ReportFiles(pdf,excel,output.getvalue(),fingerprint)
