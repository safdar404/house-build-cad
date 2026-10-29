import io,json,zipfile
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from schemas import Stage
from tools.cost import dec
from tools.model_3d import plotly_figure
from tools.schedule_4d import construction_figure

def test_cost_and_labour_reconcile(demo):
    state,_=demo;cost=state.get_output(Stage.COST);labour=state.get_output(Stage.LABOUR)
    assert sum(dec(v) for v in cost.floor_allocations_pkr.values())==dec(cost.construction_pkr)
    assert sum(dec(v) for v in cost.room_allocations_pkr.values())==dec(cost.construction_pkr)
    assert sum(dec(m['amount_pkr']) for m in cost.milestone_amounts_pkr)==dec(cost.grand_total_pkr)
    assert labour.total_wages_pkr==labour.known_wages_subtotal_pkr
    assert sum(t.paid_person_days for t in labour.totals)==labour.paid_person_days

def test_package_formats_and_caches(demo):
    state,_=demo;report=state.get_output(Stage.REPORT)
    assert len(PdfReader(io.BytesIO(report.pdf_bytes)).pages)>0
    workbook=load_workbook(io.BytesIO(report.xlsx_bytes),data_only=False)
    cached=load_workbook(io.BytesIO(report.xlsx_bytes),data_only=True)
    assert workbook['BOQ']['G2'].data_type=='f'
    assert cached['BOQ']['G2'].value is not None
    with zipfile.ZipFile(io.BytesIO(report.zip_bytes)) as z:
        assert z.testzip() is None
        suffixes={name.rsplit('.',1)[-1] for name in z.namelist()}
        assert {'pdf','xlsx','png','dxf','glb','obj'}<=suffixes
        if Stage.BIM in state.outputs:assert 'ifc' in suffixes

def test_ifc_and_4d(demo):
    state,_=demo
    if Stage.BIM in state.outputs:
        import ifcopenshell
        import ifcopenshell.validate
        model=ifcopenshell.file.from_string(state.get_output(Stage.BIM).content.decode())
        log=ifcopenshell.validate.json_logger();ifcopenshell.validate.validate(model,log,express_rules=True)
        assert not any(s.get('level')=='error' for s in log.statements)
        assert len(model.by_type('IfcBuildingStorey'))==3
    model=state.get_output(Stage.MODEL).model;schedule=state.get_output(Stage.SCHEDULE)
    animation=construction_figure(model,schedule,max_frames=300)
    assert all(t.visible is True for t in animation.frames[-1].data)
    view=plotly_figure(model,visible_floor_ids=['g'],show_roof=False)
    assert all(not trace.visible for trace in view.data if trace.meta['floor_id']!='g')
